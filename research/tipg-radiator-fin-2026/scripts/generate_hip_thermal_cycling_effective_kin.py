"""Plot effective in-plane conductivity of the HIP Ti-PG fin versus thermal cycles."""

from dataclasses import dataclass
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

from settings import output_path, workbook_path


WORKBOOK = workbook_path()
SHEET_NAME = "thermal_cycling"
OUTPUT = output_path("IRIS_TiPG_HIP_effective_kin_vs_thermal_cycles.png")

SAMPLE_CYCLES = {
    "pre-baking": 0,
    "post-baking": 1,
    "1-cycle": 1,
    "2-cycle": 2,
    "5-cycle": 5,
    "10-cycle": 10,
}
FILLETED_COLOR = "#ed6504"
CHAMFERED_COLOR = "#3d82b8"


@dataclass(frozen=True)
class CycleRecord:
    test_id: str
    sample: str
    cycles: int
    kin_chamfered: float
    kin_filleted: float


def read_cycle_data() -> tuple[CycleRecord, ...]:
    """Read and validate the dedicated thermal-cycling data tab."""
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    try:
        sheet = workbook[SHEET_NAME]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        required_headers = {"Test ID", "Sample", "k_in_chamfer", "k_in_fillet"}
        if not required_headers.issubset(headers):
            raise ValueError(
                f"{SHEET_NAME!r} must contain {sorted(required_headers)}; found {headers}."
            )
        header_index = {header: index for index, header in enumerate(headers)}
        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in row):
                continue
            test_id = str(row[header_index["Test ID"]]).strip()
            sample = str(row[header_index["Sample"]]).strip()
            if sample not in SAMPLE_CYCLES:
                raise ValueError(f"Unrecognized thermal-cycling sample label: {sample!r}.")
            records.append(
                CycleRecord(
                    test_id=test_id,
                    sample=sample,
                    cycles=SAMPLE_CYCLES[sample],
                    kin_chamfered=float(row[header_index["k_in_chamfer"]]),
                    kin_filleted=float(row[header_index["k_in_fillet"]]),
                )
            )
    finally:
        workbook.close()

    expected_samples = set(SAMPLE_CYCLES)
    observed_samples = {record.sample for record in records}
    if observed_samples != expected_samples:
        raise ValueError(
            "The thermal-cycling tab must contain exactly the expected states: "
            f"{sorted(expected_samples)}. Found {sorted(observed_samples)}."
        )
    return tuple(sorted(records, key=lambda record: (record.cycles, record.sample != "post-baking")))


def make_plot(records: tuple[CycleRecord, ...]) -> None:
    regular = [record for record in records if record.sample != "post-baking"]
    post_baking = next(record for record in records if record.sample == "post-baking")

    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 18,
            "axes.labelsize": 20,
            "xtick.labelsize": 18,
            "ytick.labelsize": 18,
            "legend.fontsize": 15,
            "mathtext.fontset": "custom",
            "mathtext.rm": "Arial",
            "mathtext.it": "Arial:italic",
        }
    )
    fig, ax = plt.subplots(figsize=(8.05, 6.04), dpi=300)
    cycles = np.array([record.cycles for record in regular], dtype=float)
    filleted = np.array([record.kin_filleted for record in regular], dtype=float)
    chamfered = np.array([record.kin_chamfered for record in regular], dtype=float)

    ax.plot(
        cycles,
        filleted,
        color=FILLETED_COLOR,
        marker="o",
        markersize=10,
        markerfacecolor="white",
        markeredgewidth=2.5,
        linestyle="none",
        label="Filleted side, cycles",
    )
    ax.plot(
        cycles,
        chamfered,
        color=CHAMFERED_COLOR,
        marker="s",
        markersize=9,
        markerfacecolor="white",
        markeredgewidth=2.5,
        linestyle="none",
        label="Chamfered side, cycles",
    )
    ax.plot(
        post_baking.cycles,
        post_baking.kin_filleted,
        linestyle="none",
        color=FILLETED_COLOR,
        marker="o",
        markersize=11,
        markeredgewidth=2,
        label="Filleted side, baked",
    )
    ax.plot(
        post_baking.cycles,
        post_baking.kin_chamfered,
        linestyle="none",
        color=CHAMFERED_COLOR,
        marker="s",
        markersize=10,
        markeredgewidth=2,
        label="Chamfered side, baked",
    )

    ax.set_xlabel("Number of Thermal Cycles")
    ax.set_ylabel(r"Effective $k_{\mathrm{in}}$ (W/(m$\cdot$K))")
    ax.set_xlim(-0.4, 10.6)
    ax.set_xticks((0, 1, 2, 5, 10))
    ax.set_ylim(0, 40)
    ax.grid(True, color="#b0b0b0", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(
        loc="lower right",
        frameon=True,
        handlelength=1.0,
        handletextpad=0.45,
        borderpad=0.35,
        labelspacing=0.3,
    )
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.0)
    fig.subplots_adjust(left=0.13, right=0.90, bottom=0.13, top=0.96)
    fig.savefig(OUTPUT, dpi=300)
    plt.close(fig)


if __name__ == "__main__":
    make_plot(read_cycle_data())
