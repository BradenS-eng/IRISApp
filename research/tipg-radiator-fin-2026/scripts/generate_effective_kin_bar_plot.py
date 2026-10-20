"""Generate the effective in-plane conductivity summary plot from the project workbook."""

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

from settings import output_path, workbook_path


WORKBOOK = workbook_path()
FIGURE_OUTPUT = output_path(
    "IRIS_effective_kin_filleted_chamfered_mean_std.png"
)
SHEET_NAME = "Sheet1"
CATEGORIES = ("Ti", "Ti-PG,\nsnap-on", "Ti-PG,\nHIP", "Ti-PG, HIP,\nbaked")
FILLETED_COLOR = "#ed6504"
CHAMFERED_COLOR = "#3d82b8"


def read_values() -> tuple[np.ndarray, np.ndarray]:
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    sheet = workbook[SHEET_NAME]

    filleted = np.array(
        [[sheet.cell(row, col).value for row in range(12, 18)] for col in range(2, 6)],
        dtype=float,
    )
    chamfered = np.array(
        [[sheet.cell(row, col).value for row in range(23, 29)] for col in range(2, 6)],
        dtype=float,
    )

    # The labeled supplemental cells record the Test 137 replacement for the baked 30 C case.
    if sheet["K12"].value != 137 or sheet["K23"].value != 137:
        raise ValueError("The workbook does not contain the expected labeled Test 137 values.")
    filleted[3, 0] = float(sheet["L12"].value)
    chamfered[3, 0] = float(sheet["L23"].value)
    return filleted, chamfered


def make_plot(filleted: np.ndarray, chamfered: np.ndarray) -> None:
    means_filleted = np.mean(filleted, axis=1)
    means_chamfered = np.mean(chamfered, axis=1)
    std_filleted = np.std(filleted, axis=1, ddof=1)
    std_chamfered = np.std(chamfered, axis=1, ddof=1)

    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 18,
            "axes.labelsize": 20,
            "xtick.labelsize": 18,
            "ytick.labelsize": 18,
            "legend.fontsize": 16,
            "mathtext.fontset": "custom",
            "mathtext.rm": "Arial",
            "mathtext.it": "Arial:italic",
        }
    )
    fig, ax = plt.subplots(figsize=(8, 6), dpi=300)
    positions = np.arange(len(CATEGORIES))
    width = 0.34
    error_style = {"ecolor": "black", "elinewidth": 1.5, "capsize": 5, "capthick": 1.5}

    ax.bar(
        positions - width / 2,
        means_filleted,
        width,
        yerr=std_filleted,
        color=FILLETED_COLOR,
        edgecolor="black",
        linewidth=1.0,
        label="Filleted side",
        error_kw=error_style,
    )
    ax.bar(
        positions + width / 2,
        means_chamfered,
        width,
        yerr=std_chamfered,
        color=CHAMFERED_COLOR,
        edgecolor="black",
        linewidth=1.0,
        label="Chamfered side",
        error_kw=error_style,
    )

    ax.set_xticks(positions, CATEGORIES)
    ax.set_ylabel(r"Effective $k_{\mathrm{in}}$ (W/(m$\cdot$K))")
    ax.set_ylim(0, 48)
    ax.grid(axis="y", color="#b0b0b0", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(loc="upper left", frameon=True)
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color("black")
        spine.set_linewidth(1.0)
    fig.tight_layout()
    fig.savefig(FIGURE_OUTPUT, dpi=300)
    plt.close(fig)

    for category, mean_f, std_f, mean_c, std_c in zip(
        CATEGORIES, means_filleted, std_filleted, means_chamfered, std_chamfered
    ):
        print(
            f"{category.replace(chr(10), ' ')}: "
            f"filleted = {mean_f:.3f} +/- {std_f:.3f}; "
            f"chamfered = {mean_c:.3f} +/- {std_c:.3f} W/(m K)"
        )


if __name__ == "__main__":
    make_plot(*read_values())
